#!/usr/bin/env python3
"""Build 12 plan docx, 8 contract docx, and manifest.json."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

ROOT = Path("/Users/nixbi/Documents/Alfa-хакатон")
PLANS = ROOT / "plans"
CONTRACTS = ROOT / "contracts"

DASH = " — "
ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def set_run_font(run, name, size, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    lang.set(qn("w:bidi"), "ru-RU")


def apply_normal(doc, font, size):
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(size)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), font)


def new_doc(font="Times New Roman", size=12, margins=(2.0, 2.0, 2.5, 2.0)):
    doc = Document()
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)
    apply_normal(doc, font, size)
    top, bottom, left, right = margins
    for section in doc.sections:
        section.top_margin = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left)
        section.right_margin = Cm(right)
        section.page_width = Twips(11906)
        section.page_height = Twips(16838)
    return doc


def add_p(
    doc,
    text,
    font,
    size=12,
    bold=False,
    italic=False,
    align="left",
    indent=None,
    before=0,
    after=8,
    line=None,
):
    p = doc.add_paragraph()
    p.alignment = ALIGN[align]
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if indent:
        fmt.first_line_indent = Cm(indent)
    if line:
        fmt.line_spacing = line
    run = p.add_run(text)
    set_run_font(run, font, size, bold, italic)
    return p


def add_h(doc, text, font, size, level=1, before=12, after=6):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    for run in p.runs:
        set_run_font(run, font, size, bold=True, italic=False)
    return p


def add_runs(doc, parts, font, size=12, align="left", before=0, after=8, indent=None):
    p = doc.add_paragraph()
    p.alignment = ALIGN[align]
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if indent:
        fmt.first_line_indent = Cm(indent)
    for text, bold, italic in parts:
        run = p.add_run(text)
        set_run_font(run, font, size, bold, italic)
    return p


def extract(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def grouped(n):
    return f"{n:,}".replace(",", " ")


# ---------------------------------------------------------------------------
# Plans
# ---------------------------------------------------------------------------

def build_plan_01():
    font = "Times New Roman"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Служебная записка"
    doc.core_properties.author = "Инна Корсакова"
    add_h(doc, "Служебная записка", font, 16, level=1, before=0)
    add_h(doc, "Суть", font, 13, level=2)
    context = (
        "Мастерская «Резец и лак» делает простую деревянную мебель и продаёт её гостиницам "
        "«Светлая пристань»: заработок складывается из счетов за партии стульев и столов для этих гостиниц."
    )
    add_p(doc, context, font, align="justify", indent=1.25, after=6)
    add_h(doc, "Показатели месяца", font, 13, level=2)
    revenue = "Месячная выручка составляет 400 000 руб. в месяц."
    costs = "Ежемесячные расходы равны 0,55 от выручки, то есть 220 000 руб. в месяц."
    own = "Собственные средства составляют 100 000 руб."
    months = "До первых денежных поступлений осталось 2 месяца."
    for line in (revenue, costs, own, months):
        add_p(doc, line, font, align="justify", indent=1.25, after=4)
    add_h(doc, "На что нужны деньги", font, 13, level=2)
    use = "Деньги нужны на оборудование."
    add_p(doc, use, font, align="justify", indent=1.25, after=4)
    return doc, {
        "revenue_monthly_rub": revenue,
        "costs_monthly_rub": costs,
        "own_funds_rub": own,
        "months_to_first_cashflow": months,
        "use_of_funds": use,
        "context_specific": context,
    }


def build_plan_02():
    font = "Georgia"
    doc = new_doc(font, 12, margins=(2.2, 2.2, 3.0, 2.5))
    doc.core_properties.title = "Рассказ о кофейне"
    doc.core_properties.author = "Леонид Пашин"
    text = (
        "Мы в кофейне «Пенка и корица» варим напитки и печём булочки для сотрудников бюро "
        "«Тихий чертёж», заработок идёт со счетов за их ежедневные завтраки. "
        "Месячная выручка составляет 80 000 руб. в месяц, а расходы забирают 0,90 от выручки, "
        "это 72 000 руб. в месяц. Поступления уже есть, до них 0 месяцев. "
        "Деньги нужны на оборотные средства, чтобы закрывать закуп кофе до оплаты от бюро."
    )
    add_p(doc, text, font, align="justify", indent=1.25, after=0, line=1.15)
    return doc, {
        "revenue_monthly_rub": "Месячная выручка составляет 80 000 руб. в месяц",
        "costs_monthly_rub": "расходы забирают 0,90 от выручки, это 72 000 руб. в месяц",
        "months_to_first_cashflow": "Поступления уже есть, до них 0 месяцев.",
        "use_of_funds": "Деньги нужны на оборотные средства",
        "context_specific": (
            "Мы в кофейне «Пенка и корица» варим напитки и печём булочки для сотрудников бюро "
            "«Тихий чертёж», заработок идёт со счетов за их ежедневные завтраки."
        ),
    }


def build_plan_03():
    font = "Menlo"
    doc = new_doc(font, 12, margins=(1.8, 1.8, 2.0, 2.0))
    doc.core_properties.title = "Обрывки"
    doc.core_properties.author = "Савва Иртышев"
    lines = {
        "context_specific": "городское дело, обычный формат",
        "revenue_monthly_rub": "выручка 1 500 000 руб. в месяц",
        "own_funds_rub": "свои деньги 700 000 руб.",
        "months_to_first_cashflow": "до потока 6 месяцев",
        "use_of_funds": "деньги на запуск",
    }
    order = [
        "context_specific",
        "revenue_monthly_rub",
        "own_funds_rub",
        "months_to_first_cashflow",
        "use_of_funds",
    ]
    for key in order:
        add_p(doc, lines[key], font, size=12, after=0, before=0)
    return doc, lines


def build_plan_04():
    font = "Arial"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Письмо партнёру"
    doc.core_properties.author = "Полина Верховая"
    add_p(doc, "Ладогорье, март", font, align="right", after=0)
    add_p(doc, "Кому: Глебу Ракитину", font, align="right", after=0)
    add_p(doc, "От: Полины Верховой", font, align="right", after=12)
    add_p(doc, "Глеб, здравствуй.", font, after=8)
    context = (
        "Пишу тебе как партнёру. Цех «Серый фургон» возит узлы для сборки компании "
        "«Ковш и перец», заработок берём с оплаты этих рейсов."
    )
    add_p(doc, context, font, align="justify", after=8)
    use = "Деньги нужны на оборудование."
    own = "Свои деньги, которые могу вложить, — 100 000 руб."
    months = "До первых поступлений по новому маршруту 10 месяцев."
    add_p(doc, use, font, align="justify", after=8)
    add_p(doc, own, font, align="justify", after=8)
    add_p(doc, months, font, align="justify", after=8)
    add_p(doc, "Если откликнешься, созвонимся и обсудим поставку узлов.", font, align="justify", after=12)
    add_p(doc, "Полина Верховая", font, after=0)
    return doc, {
        "own_funds_rub": own,
        "months_to_first_cashflow": months,
        "use_of_funds": use,
        "context_specific": (
            "Цех «Серый фургон» возит узлы для сборки компании «Ковш и перец», "
            "заработок берём с оплаты этих рейсов."
        ),
    }


def build_plan_05():
    font = "Verdana"
    doc = new_doc(font, 11)
    doc.core_properties.title = "Перечень показателей"
    doc.core_properties.author = "Кира Луговая"
    rows = [
        "Документ — перечень показателей",
        "Выручка — 400 000 руб. в месяц",
        "Расходы — 1,25 от выручки, 500 000 руб. в месяц",
        "Назначение — оборотные средства",
        "Описание — обычная городская деятельность в целом",
    ]
    # The source uses an em dash. Keep the file text identical to DASH.
    rows = [
        f"Документ{DASH}перечень показателей",
        f"Выручка{DASH}400 000 руб. в месяц",
        f"Расходы{DASH}1,25 от выручки, 500 000 руб. в месяц",
        f"Назначение{DASH}оборотные средства",
        f"Описание{DASH}обычная городская деятельность в целом",
    ]
    for row in rows:
        add_p(doc, row, font, size=11, after=2, before=2)
    return doc, {
        "revenue_monthly_rub": rows[1],
        "costs_monthly_rub": rows[2],
        "use_of_funds": rows[3],
        "context_specific": rows[4],
    }


def build_plan_06():
    font = "Palatino"
    doc = new_doc(font, 12, margins=(2.0, 2.0, 2.8, 2.4))
    doc.core_properties.title = "Пасека за оврагом"
    doc.core_properties.author = "Арсений Мохов"
    story = [
        "Пасека за оврагом",
        "Арсений Мохов держит пасеку «Улей за оврагом» за ручьём, у окраины Ладогорья. Тропа туда идёт через ольшаник, и весной она раскисает так, что фляги проще катить на санках, чем нести в руках.",
        "Зимой он сидит над журналом осмотров, чинит рамки и сушит воск. Летом живёт у ульев почти безвыездно: смотрит расплод, подставляет магазины и отгоняет муравьёв от поилки.",
        "Мёд и вощину у него забирает пекарня «Булка до полудня», а заработок складывается из оплаты этих партий: пекарня печёт пряники и платит по счёту после каждой сдачи.",
        "С пекарней они знакомы давно. Хозяйка приезжает сама, пробует мёд ложкой прямо с фляги и только потом кивает грузчику. Если вкус уходит в горчинку, партию оставляют на пасеке и переливают в другую тару.",
        "Арсений не любит торопить отгрузку. Он даёт мёду отстояться, снимает пену и только затем закрывает флягу. Вощину складывает отдельно, чтобы она не впитала запах дымаря.",
        "В дождливые недели дорога перекрыта ручьём, и сдачу переносят, пока вода не спадёт. Пекарня ждёт: пряники в эти дни пекут на старом запасе, а новую варку ставят, когда фляги всё-таки доезжают.",
        "На пасеке помогают племянница и сосед. Племянница следит за записями в тетради, сосед возит фляги на тележке до сухой грунтовки, где их забирает машина пекарни. Разговор обычно короткий: сверились по названиям сортов, поставили подписи в тетради и разъехались.",
        "Арсений описывает дело спокойно, без громких обещаний. Ему важно, чтобы пекарня и дальше брала именно его мёд, а не переходила на чужой, и чтобы оплата приходила после каждой сдачи, как условились устно ещё при первой варке пряников.",
        "Так и устроена работа: пасека сдаёт продукт пекарне, пекарня платит за партии, а сезон крутится от весенней чистки рамок до осеннего розлива.",
        "В тихие вечера он перечитывает тетрадь и раскладывает ярлыки по сортам. Липовый мёд идёт в пряники, гречишный оставляют для отдельной полки, а вощина уезжает вместе с ближайшей флягой, чтобы не копилась у окна.",
        "Хозяйка пекарни иногда остаётся на чай. Говорят о том, какая партия вышла мягче, и о том, что дорогу надо подсыпать щебнем, иначе тележка снова увязнет. Никаких бумаг сверх тетради они в эти вечера не заводят.",
        "Сосед обещает, что ручей когда-нибудь перекроют мостками. Пока мостков не поставили, Арсений закладывает лишний час на выезд и не обещает пекарне день, если с ночи шёл дождь.",
        "Племянница подписывает ярлыки аккуратно, буква к букве, и убирает тетрадь в жестяную коробку. Коробка стоит на полке над весами, рядом с запасными рамками и мотком шпагата.",
    ]
    numbers = (
        "Дальше в этом рассказе появляются суммы. Месячная выручка составляет 1 500 000 руб. в месяц. "
        "Расходы составляют 0,55 от выручки, то есть 825 000 руб. в месяц. "
        "Собственные средства составляют 700 000 руб. До первых денежных поступлений в новом сезоне 6 месяцев."
    )
    add_p(doc, story[0], font, size=18, italic=True, align="center", after=12)
    for para in story[1:]:
        add_p(doc, para, font, align="justify", indent=1.25, after=8, line=1.15)
    add_p(doc, numbers, font, align="justify", indent=1.25, before=12, after=0, line=1.15)
    return doc, {
        "revenue_monthly_rub": "Месячная выручка составляет 1 500 000 руб. в месяц.",
        "costs_monthly_rub": "Расходы составляют 0,55 от выручки, то есть 825 000 руб. в месяц.",
        "own_funds_rub": "Собственные средства составляют 700 000 руб.",
        "months_to_first_cashflow": "До первых денежных поступлений в новом сезоне 6 месяцев.",
        "context_specific": story[3],
    }


def build_plan_07():
    font = "Times"
    doc = new_doc(font, 12, margins=(1.7, 1.7, 2.2, 2.0))
    doc.core_properties.title = "Справка бухгалтера"
    doc.core_properties.author = "Эмма Шилова"
    add_p(doc, "СПРАВКА", font, size=16, bold=True, align="center", after=2)
    add_p(doc, "для внутреннего учёта", font, size=11, align="center", after=10)
    context = "Настоящая справка составлена по обычной городской деятельности, описанной в целом."
    add_p(doc, context, font, align="justify", after=8)
    add_p(
        doc,
        "Нижеследующее изложено по сведениям, имеющимся на дату подготовки, и подлежит переносу в учётные регистры.",
        font,
        align="justify",
        after=8,
    )
    revenue = "1. Ежемесячная выручка составляет 80 000 руб."
    costs = "2. Расходы исчислены как 0,90 от выручки и составляют 72 000 руб. в месяц."
    own = "3. Собственные средства отражены в сумме 100 000 руб."
    use = "4. Назначение средств определено так: запуск."
    for line in (revenue, costs, own, use):
        add_p(doc, line, font, align="justify", after=4)
    add_p(doc, "Справку составила бухгалтер Эмма Шилова.", font, before=10, after=0)
    return doc, {
        "revenue_monthly_rub": revenue,
        "costs_monthly_rub": costs,
        "own_funds_rub": own,
        "use_of_funds": use,
        "context_specific": context,
    }


def build_plan_08():
    font = "Gill Sans"
    doc = new_doc(font, 14, margins=(3.0, 3.0, 3.0, 3.0))
    doc.core_properties.title = "Короткая записка"
    doc.core_properties.author = "Мирон Янцев"
    text = (
        "Мастерская «Ключ у порога» чинит технику жильцам дома «Кедровый» и зарабатывает на этих выездах. "
        "Деньги нужны на оборудование. Выручка 400 000 руб. в месяц. До первых поступлений 0 месяцев."
    )
    add_p(doc, text, font, size=14, align="center", before=36, after=0, line=1.3)
    return doc, {
        "revenue_monthly_rub": "Выручка 400 000 руб. в месяц.",
        "months_to_first_cashflow": "До первых поступлений 0 месяцев.",
        "use_of_funds": "Деньги нужны на оборудование.",
        "context_specific": (
            "Мастерская «Ключ у порога» чинит технику жильцам дома «Кедровый» и зарабатывает на этих выездах."
        ),
    }


def build_plan_09():
    font = "Trebuchet MS"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Пропись и цифры"
    doc.core_properties.author = "Нина Соколова"
    context = (
        "Служба «Мятный ковёр» убирает офисы бюро «Ровный край» и зарабатывает на ежемесячных счетах за эту уборку."
    )
    use = "Деньги нужны на оборотные средства."
    revenue = "Месячная выручка составляет один миллион пятьсот тысяч руб. в месяц."
    costs = "Расходы равны 1,25 от выручки, это 1 875 000 руб. в месяц."
    own = "Собственные средства составляют семьсот тысяч руб."
    months = "До первого денежного потока 10 месяцев."
    add_p(doc, context, font, align="justify", after=8)
    add_p(doc, use, font, align="justify", after=8)
    add_p(doc, revenue, font, italic=True, align="justify", after=8)
    add_runs(
        doc,
        [
            ("Расходы равны 1,25 от выручки, это ", False, False),
            ("1 875 000", True, False),
            (" руб. в месяц.", False, False),
        ],
        font,
        align="justify",
    )
    add_p(doc, own, font, italic=True, align="justify", after=8)
    add_p(doc, months, font, align="justify", after=0)
    return doc, {
        "revenue_monthly_rub": revenue,
        "costs_monthly_rub": costs,
        "own_funds_rub": own,
        "months_to_first_cashflow": months,
        "use_of_funds": use,
        "context_specific": context,
    }


def build_plan_10():
    font = "Avenir"
    doc = new_doc(font, 13, margins=(2.0, 2.0, 2.2, 2.2))
    doc.core_properties.title = "Как есть"
    doc.core_properties.author = "Глеб Ракитин"
    text = (
        "Короче, дело у нас городское и обычное вобще, я так это вижу. "
        "До первых денег ещё 2 месяца. На запуск денги и нужнны, остально сам допишу када время будет."
    )
    add_p(doc, text, font, size=13, align="left", after=0, line=1.2)
    return doc, {
        "months_to_first_cashflow": "До первых денег ещё 2 месяца.",
        "use_of_funds": "На запуск денги и нужнны, остально сам допишу када время будет.",
        "context_specific": "Короче, дело у нас городское и обычное вобще, я так это вижу.",
    }


def build_plan_11():
    font = "Helvetica"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Сначала смысл"
    doc.core_properties.author = "Марк Ельцов"
    context = (
        "Возим фурнитуру цеху «Доска и клей»: забираем её со склада и зарабатываем на плате за эти рейсы."
    )
    use = "Кредит нужен на оборотные средства."
    revenue = "Месячная выручка составляет 80 000 руб. в месяц."
    costs = "Расходы составляют 0,55 от выручки, то есть 44 000 руб. в месяц."
    own = "Собственные средства составляют 100 000 руб."
    add_p(doc, context, font, align="justify", after=8)
    add_p(doc, use, font, align="justify", after=8)
    add_p(doc, revenue, font, align="justify", before=18, after=6)
    add_p(doc, costs, font, align="justify", after=6)
    add_p(doc, own, font, align="justify", after=0)
    return doc, {
        "revenue_monthly_rub": revenue,
        "costs_monthly_rub": costs,
        "own_funds_rub": own,
        "use_of_funds": use,
        "context_specific": context,
    }


def build_plan_12():
    font = "Courier New"
    doc = new_doc(font, 12, margins=(2.0, 2.0, 2.0, 2.0))
    doc.core_properties.title = "Строки показателей"
    doc.core_properties.author = "Эмма Шилова"
    lines = [
        "показатели за месяц",
        "выручка | 400000",
        "расходы | 360000",
        "доля расходов | 0.90",
        "месяцы до потока | 6",
    ]
    for line in lines:
        add_p(doc, line, font, size=12, after=0, before=0)
    return doc, {
        "revenue_monthly_rub": "выручка | 400000",
        "costs_monthly_rub": "расходы | 360000",
        "months_to_first_cashflow": "месяцы до потока | 6",
    }


# ---------------------------------------------------------------------------
# Contracts
# ---------------------------------------------------------------------------

def build_contract_01():
    font = "Times New Roman"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Договор поставки"
    doc.core_properties.author = "Нина Соколова"
    add_p(doc, "ДОГОВОР ПОСТАВКИ", font, size=16, bold=True, align="center", after=12)
    quote = (
        "ООО «Синий ящик», именуемое Поставщик, и ООО «Свёрток и шпагат», именуемое Покупатель, "
        "заключили договор поставки."
    )
    add_p(doc, quote, font, align="justify", indent=1.25, after=8)
    add_p(
        doc,
        "Поставщик передаёт Покупателю партию хозяйственной тары, а Покупатель принимает её по накладной.",
        font,
        align="justify",
        indent=1.25,
        after=8,
    )
    add_p(
        doc,
        "Общая сумма поставки составляет 186 400 руб.",
        font,
        align="justify",
        indent=1.25,
        after=8,
    )
    add_p(
        doc,
        "Качество тары стороны смотрят при приёмке и фиксируют в накладной.",
        font,
        align="justify",
        indent=1.25,
        after=12,
    )
    add_p(doc, "Поставщик: Нина Соколова", font, after=2)
    add_p(doc, "Покупатель: Марк Ельцов", font, after=0)
    return doc, {"has_contract": quote}


def build_contract_02():
    font = "Calibri"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Договор оказания услуг"
    doc.core_properties.author = "Кира Луговая"
    add_p(doc, "ДОГОВОР ОКАЗАНИЯ УСЛУГ", font, size=16, bold=True, align="center", after=12)
    quote = "Настоящий договор оказания услуг заключён между ООО «Мятный ковёр» и ООО «Тихий борт»."
    add_p(doc, quote, font, align="justify", after=8)
    add_p(doc, "Исполнитель убирает складские помещения заказчика.", font, align="justify", after=8)
    add_p(doc, "Сумма услуг составляет 94 800 руб.", font, align="justify", after=4)
    add_p(doc, "Срок оказания услуг — 45 календарных дней с даты подписания.", font, align="justify", after=10)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    cells = (
        ("Показатель", "Значение"),
        ("Сумма", "94 800 руб."),
        ("Срок", "45 календарных дней"),
    )
    for i, (left, right) in enumerate(cells):
        table.rows[i].cells[0].text = left
        table.rows[i].cells[1].text = right
        for cell in table.rows[i].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, font, 11, bold=(i == 0))
    return doc, {"has_contract": quote}


def build_contract_03():
    font = "Georgia"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Договор аренды помещения"
    doc.core_properties.author = "Марк Ельцов"
    add_p(doc, "ДОГОВОР АРЕНДЫ ПОМЕЩЕНИЯ", font, size=16, bold=True, align="center", after=12)
    quote = (
        "ООО «Прожектор в чехле» арендует у ООО «Лампа на вечер» нежилое помещение "
        "по адресу: г. Ладогорье, пер. Кедровый, д. 7."
    )
    add_p(doc, quote, font, align="justify", indent=1.25, after=8)
    add_p(
        doc,
        "Помещение используется под хранение сценического света.",
        font,
        align="justify",
        indent=1.25,
        after=8,
    )
    add_p(
        doc,
        "Арендная плата по этому договору составляет 73 500 руб.",
        font,
        align="justify",
        indent=1.25,
        after=8,
    )
    add_p(doc, "Арендодатель: Лампа на вечер", font, before=8, after=2)
    add_p(doc, "Арендатор: Прожектор в чехле", font, after=0)
    return doc, {"has_contract": quote}


def build_contract_04():
    font = "Arial"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Предварительный заказ"
    doc.core_properties.author = "Полина Верховая"
    add_p(doc, "ПРЕДВАРИТЕЛЬНЫЙ ЗАКАЗ", font, size=16, bold=True, align="center", after=12)
    quote = (
        "ООО «Ковш и перец» направляет ООО «Стеклянный край» предварительный заказ "
        "на стеклянные бутылки для соусов."
    )
    add_p(doc, quote, font, align="justify", after=8)
    add_p(
        doc,
        "Заказ фиксирует намерение сторон и подлежит уточнению по объёму перед отгрузкой.",
        font,
        align="justify",
        after=8,
    )
    add_p(doc, "Ориентир оплаты за партию — 128 900 руб.", font, align="justify", after=0)
    return doc, {"has_contract": quote}


def build_contract_05():
    font = "Times New Roman"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Рамочный договор"
    doc.core_properties.author = "Савва Иртышев"
    add_p(doc, "РАМОЧНЫЙ ДОГОВОР", font, size=16, bold=True, align="center", after=12)
    add_p(
        doc,
        "ООО «Серый фургон» и ООО «Доска и клей» заключают рамочный договор о перевозках.",
        font,
        align="justify",
        indent=1.25,
        after=8,
    )
    quote = (
        "Цена перевозок в настоящем документе твёрдой суммой не закрепляется "
        "и будет указана в отдельных заявках."
    )
    add_p(doc, quote, font, align="justify", indent=1.25, after=8)
    add_p(
        doc,
        "Предмет ограничивается согласованием порядка будущих рейсов.",
        font,
        align="justify",
        indent=1.25,
        after=8,
    )
    add_p(
        doc,
        "Заявки стороны оформляют по мере появления груза.",
        font,
        align="justify",
        indent=1.25,
        after=0,
    )
    return doc, {"has_contract": quote}


def build_contract_06():
    font = "Georgia"
    doc = new_doc(font, 12, margins=(2.5, 2.5, 2.5, 2.5))
    doc.core_properties.title = "Обрывок листа"
    doc.core_properties.author = "Савва Иртышев"
    paragraphs = [
        "…после чего стороны сверяют комплектность тары и подписывают накладную прямо на рампе. Кладовщик относит целые коробки к стеллажу у дальней стены, а представитель второй стороны остаётся до конца сверки и смотрит, чтобы ничего не смешали с чужим грузом.",
        "Если угол коробки помят, её откладывают на отдельный поддон и описывают в акте, который составляют тут же от руки, пока водитель не закрыл тент. Споры по комплектности стараются закрыть в тот же визит, иначе машина уедет, а спор останется на словах.",
        "Погрузчик работает только под навесом: во дворе сыро, и картон быстро набирает влагу, отчего крышки коробит. Ответственный со стороны получателя отмечает в журнале фамилию водителя и время въезда, без каких-либо денежных пометок.",
        "Когда сверка затягивается, сторож пускает людей в тёплую комнату и наливает кипяток из бака. Разговор там тихий, про ремни и про то, какая партия поедет следующей, а лист на этом месте обрывается, и фраза остаётся незаконченной: кладовщик повернулся к стеллажу и только начал говорить, что крайняя пачка…",
    ]
    for para in paragraphs:
        add_p(doc, para, font, italic=True, align="justify", indent=1.0, after=8, line=1.15)
    quote = "Ответственный со стороны получателя отмечает в журнале фамилию водителя и время въезда, без каких-либо денежных пометок."
    return doc, {"has_contract": quote}


def build_contract_07():
    font = "Calibri"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Письмо"
    doc.core_properties.author = "Кирилл Лемешев"
    add_p(doc, "Ольга, здравствуйте.", font, after=8)
    add_p(
        doc,
        "Меня зовут Кирилл Лемешев, я из мастерской «Ключ у порога».",
        font,
        after=8,
    )
    quote = "Хочу сказать прямо: давайте работать."
    add_p(doc, quote, font, after=8)
    add_p(
        doc,
        "Мне близка ваша аккуратность. Буду рад созвониться, когда вам будет удобно.",
        font,
        after=12,
    )
    add_p(doc, "Кирилл Лемешев", font, after=0)
    return doc, {"has_contract": quote}


def build_contract_08():
    font = "Courier New"
    doc = new_doc(font, 12)
    doc.core_properties.title = "Бланк договора"
    doc.core_properties.author = "Форма"
    add_p(doc, "БЛАНК ДОГОВОРА", font, size=16, bold=True, align="center", after=12)
    quote = "Форма не заполнена."
    add_p(doc, quote, font, align="center", after=12)
    for line in (
        "Сторона А: ________",
        "Сторона Б: ________",
        "Предмет: ________",
        "Сумма: ________",
    ):
        add_p(doc, line, font, size=12, after=6)
    return doc, {"has_contract": quote}


PLANS_SPEC = [
    ("01_деловая_записка.docx", "01", build_plan_01, {
        "revenue_monthly_rub": 400000,
        "costs_monthly_rub": 220000,
        "own_funds_rub": 100000,
        "months_to_first_cashflow": 2,
        "use_of_funds": "оборудование",
        "context_specific": True,
        "has_contract": None,
    }),
    ("02_монолог.docx", "02", build_plan_02, {
        "revenue_monthly_rub": 80000,
        "costs_monthly_rub": 72000,
        "own_funds_rub": None,
        "months_to_first_cashflow": 0,
        "use_of_funds": "оборотные средства",
        "context_specific": True,
        "has_contract": None,
    }),
    ("03_обрывки.docx", "03", build_plan_03, {
        "revenue_monthly_rub": 1500000,
        "costs_monthly_rub": None,
        "own_funds_rub": 700000,
        "months_to_first_cashflow": 6,
        "use_of_funds": "запуск",
        "context_specific": False,
        "has_contract": None,
    }),
    ("04_письмо_партнёру.docx", "04", build_plan_04, {
        "revenue_monthly_rub": None,
        "costs_monthly_rub": None,
        "own_funds_rub": 100000,
        "months_to_first_cashflow": 10,
        "use_of_funds": "оборудование",
        "context_specific": True,
        "has_contract": None,
    }),
    ("05_показатели.docx", "05", build_plan_05, {
        "revenue_monthly_rub": 400000,
        "costs_monthly_rub": 500000,
        "own_funds_rub": None,
        "months_to_first_cashflow": None,
        "use_of_funds": "оборотные средства",
        "context_specific": False,
        "has_contract": None,
    }),
    ("06_рассказ.docx", "06", build_plan_06, {
        "revenue_monthly_rub": 1500000,
        "costs_monthly_rub": 825000,
        "own_funds_rub": 700000,
        "months_to_first_cashflow": 6,
        "use_of_funds": None,
        "context_specific": True,
        "has_contract": None,
    }),
    ("07_канцелярит.docx", "07", build_plan_07, {
        "revenue_monthly_rub": 80000,
        "costs_monthly_rub": 72000,
        "own_funds_rub": 100000,
        "months_to_first_cashflow": None,
        "use_of_funds": "запуск",
        "context_specific": False,
        "has_contract": None,
    }),
    ("08_короткий.docx", "08", build_plan_08, {
        "revenue_monthly_rub": 400000,
        "costs_monthly_rub": None,
        "own_funds_rub": None,
        "months_to_first_cashflow": 0,
        "use_of_funds": "оборудование",
        "context_specific": True,
        "has_contract": None,
    }),
    ("09_пропись_и_цифры.docx", "09", build_plan_09, {
        "revenue_monthly_rub": 1500000,
        "costs_monthly_rub": 1875000,
        "own_funds_rub": 700000,
        "months_to_first_cashflow": 10,
        "use_of_funds": "оборотные средства",
        "context_specific": True,
        "has_contract": None,
    }),
    ("10_разговорный.docx", "10", build_plan_10, {
        "revenue_monthly_rub": None,
        "costs_monthly_rub": None,
        "own_funds_rub": None,
        "months_to_first_cashflow": 2,
        "use_of_funds": "запуск",
        "context_specific": False,
        "has_contract": None,
    }),
    ("11_обратный_порядок.docx", "11", build_plan_11, {
        "revenue_monthly_rub": 80000,
        "costs_monthly_rub": 44000,
        "own_funds_rub": 100000,
        "months_to_first_cashflow": None,
        "use_of_funds": "оборотные средства",
        "context_specific": True,
        "has_contract": None,
    }),
    ("12_строки.docx", "12", build_plan_12, {
        "revenue_monthly_rub": 400000,
        "costs_monthly_rub": 360000,
        "own_funds_rub": None,
        "months_to_first_cashflow": 6,
        "use_of_funds": None,
        "context_specific": None,
        "has_contract": None,
    }),
]

CONTRACTS_SPEC = [
    ("01_поставка.docx", "01", build_contract_01, True),
    ("02_услуги.docx", "02", build_contract_02, True),
    ("03_аренда.docx", "03", build_contract_03, True),
    ("04_предварительный_заказ.docx", "04", build_contract_04, True),
    ("05_рамочный.docx", "05", build_contract_05, False),
    ("06_обрывок.docx", "06", build_contract_06, False),
    ("07_письмо.docx", "07", build_contract_07, False),
    ("08_бланк.docx", "08", build_contract_08, False),
]

FACT_KEYS = [
    "revenue_monthly_rub",
    "costs_monthly_rub",
    "own_funds_rub",
    "months_to_first_cashflow",
    "use_of_funds",
    "context_specific",
    "has_contract",
]

BANNED = [
    "балл",
    "рейтинг",
    "займ",
    "заём",
    "заем",
    "скоринг",
    "подтвержд",
    "флаг",
    "одобр",
    "отказ",
]


def empty_contract_facts(has_contract):
    facts = {key: None for key in FACT_KEYS}
    facts["has_contract"] = has_contract
    return facts


def word_count(text):
    return len(re.findall(r"\S+", text))


def assert_quote(text, quote, label):
    if quote not in text:
        raise SystemExit(f"QUOTE MISSING in {label}: {quote}")


def validate_plan(path, style_id, facts, quotes, text):
    low = text.lower()
    for word in BANNED:
        if word in low:
            raise SystemExit(f"{path} contains banned word {word}")
    if re.search(r"(?<![\w])нет(?![\w])", low):
        raise SystemExit(f"{path} contains нет")
    if "срок" in low:
        raise SystemExit(f"{path} contains срок")
    if "договор" in low or "контракт" in low:
        raise SystemExit(f"{path} mentions a contract")
    for amount in ("300 000", "300000", "800 000", "800000", "2 000 000", "2000000"):
        if amount in text:
            raise SystemExit(f"{path} contains loan-like amount {amount}")

    for key, value in facts.items():
        if value is None:
            if key in quotes:
                raise SystemExit(f"{path} has quote for empty {key}")
        else:
            if key not in quotes:
                raise SystemExit(f"{path} missing quote for {key}")
            assert_quote(text, quotes[key], f"{path} {key}")

    if facts["use_of_funds"]:
        if facts["use_of_funds"] not in quotes["use_of_funds"]:
            raise SystemExit(f"{path} use quote does not contain the purpose")
    else:
        for purpose in ("запуск", "оборудование", "оборотные"):
            if purpose in low:
                raise SystemExit(f"{path} mentions purpose while use_of_funds is empty: {purpose}")

    if facts["revenue_monthly_rub"] is None:
        for token in ("80 000", "80000", "400 000", "400000", "1 500 000", "1500000",
                      "восемьдесят тысяч", "четыреста тысяч", "один миллион пятьсот тысяч"):
            if token in text:
                raise SystemExit(f"{path} has revenue token while revenue is empty: {token}")
    if facts["costs_monthly_rub"] is None:
        for token in ("0,55", "0,90", "1,25", "0.55", "0.90", "1.25"):
            if token in text:
                raise SystemExit(f"{path} has a cost ratio while costs are empty")
    if facts["own_funds_rub"] is None:
        for token in ("100 000", "100000", "700 000", "700000", "сто тысяч", "семьсот тысяч"):
            if token in text:
                raise SystemExit(f"{path} has own funds while empty")
    if facts["months_to_first_cashflow"] is None:
        if re.search(r"(?<!\d)(0|2|6|10)\s+месяц", text):
            raise SystemExit(f"{path} has months-to-cashflow while empty")

    if facts["context_specific"] is True:
        quote = quotes["context_specific"]
        if "«" not in quote or not re.search(r"зарабатыва|заработок|оплат", quote):
            raise SystemExit(f"{path} specific context quote is weak: {quote}")
    if facts["context_specific"] is False:
        if "«" in text:
            raise SystemExit(f"{path} general context names a party")
        if re.search(r"зарабатыва|заработок|покупател|клиент", low):
            raise SystemExit(f"{path} general context names an earning method or buyer")

    if style_id == "06":
        match = re.search(r"\d", text)
        if match is None or match.start() <= len(text) / 2:
            raise SystemExit(f"plan 06 digits not only in the second half: {match.start() if match else None} / {len(text)}")
    if style_id == "08":
        count = word_count(text)
        if count >= 90:
            raise SystemExit(f"plan 08 has {count} words")
    if style_id == "09":
        for token in ("один миллион пятьсот тысяч", "семьсот тысяч", "1 875 000"):
            if token not in text:
                raise SystemExit(f"plan 09 missing {token}")
    if style_id == "10":
        for token in ("вобще", "денги", "нужнны", "остально", "када"):
            if token not in text:
                raise SystemExit(f"plan 10 missing typo {token}")
    if style_id == "11":
        match = re.search(r"\d", text)
        ctx_end = text.index(quotes["context_specific"]) + len(quotes["context_specific"])
        use_end = text.index(quotes["use_of_funds"]) + len(quotes["use_of_funds"])
        if match is None or match.start() <= max(ctx_end, use_end):
            raise SystemExit("plan 11 numbers are not after context and purpose")
    if style_id == "12":
        if "выручка | 400000" not in text or "расходы | 360000" not in text:
            raise SystemExit("plan 12 pipe lines missing")
    if style_id == "05":
        for line in text.splitlines():
            if line and DASH not in line:
                raise SystemExit(f"plan 05 line without dash: {line}")
    if style_id == "01":
        headings = [p.style.name for p in Document(path).paragraphs if p.style and p.style.name.startswith("Heading")]
        if len(headings) < 3:
            raise SystemExit(f"plan 01 headings: {headings}")
    if style_id == "02":
        styles = [p.style.name for p in Document(path).paragraphs if p.text]
        if any(name.startswith("Heading") for name in styles):
            raise SystemExit("plan 02 has a heading")
        if len([p for p in Document(path).paragraphs if p.text]) != 1:
            raise SystemExit("plan 02 is not a single monologue")


def validate_contract(path, style_id, facts, quotes, text):
    low = text.lower()
    for word in BANNED:
        if word in low:
            raise SystemExit(f"{path} contains banned word {word}")
    for key, value in facts.items():
        if value is None:
            if key in quotes:
                raise SystemExit(f"{path} quotes empty {key}")
        else:
            assert_quote(text, quotes[key], f"{path} {key}")
    if facts["has_contract"] is True and style_id not in {"01", "02", "03", "04"}:
        raise SystemExit(style_id)
    if facts["has_contract"] is False and style_id not in {"05", "06", "07", "08"}:
        raise SystemExit(style_id)
    if style_id == "01":
        for token in ("Синий ящик", "Свёрток и шпагат", "186 400"):
            if token not in text:
                raise SystemExit(f"contract 01 missing {token}")
    if style_id == "02":
        for token in ("94 800", "Срок", "45 календарных дней"):
            if token not in text:
                raise SystemExit(f"contract 02 missing {token}")
    if style_id == "03":
        for token in ("Ладогорье", "пер. Кедровый", "д. 7", "аренд"):
            if token not in text.lower() and token not in text:
                raise SystemExit(f"contract 03 missing {token}")
    if style_id == "04":
        if "предварительный заказ" not in low:
            raise SystemExit("contract 04 is not a preliminary order")
    if style_id == "05":
        if re.search(r"\d", text):
            raise SystemExit("contract 05 has a digit")
        if "твёрдой суммой не закрепляется" not in text:
            raise SystemExit("contract 05 looks priced")
    if style_id == "06":
        if re.search(r"\d", text) or "руб" in low:
            raise SystemExit("contract 06 has an amount")
        count = word_count(text)
        if not 140 <= count <= 260:
            raise SystemExit(f"contract 06 word count {count}")
    if style_id == "07":
        if "давайте работать" not in text:
            raise SystemExit("contract 07 missing the phrase")
        if re.search(r"\d", text) or "руб" in low:
            raise SystemExit("contract 07 looks like a priced contract")
    if style_id == "08":
        if re.search(r"\d", text):
            raise SystemExit("contract 08 has a digit")
        if "________" not in text:
            raise SystemExit("contract 08 is not a blank")
        if "Синий" in text or "руб" in low:
            raise SystemExit("contract 08 is filled")


def record(file_rel, kind, style_id, facts, quotes):
    item = {"file": file_rel, "kind": kind, "style_id": style_id}
    for key in FACT_KEYS:
        item[key] = facts[key]
    item["quotes"] = {key: quotes[key] for key in FACT_KEYS if facts[key] is not None}
    return item


def main():
    PLANS.mkdir(exist_ok=True)
    CONTRACTS.mkdir(exist_ok=True)
    manifest = []

    for filename, style_id, builder, facts in PLANS_SPEC:
        doc, quotes = builder()
        path = PLANS / filename
        doc.save(path)
        if not zipfile.is_zipfile(path):
            raise SystemExit(f"not a docx: {path}")
        with zipfile.ZipFile(path) as zf:
            if "word/document.xml" not in zf.namelist():
                raise SystemExit(f"missing document.xml: {path}")
        reloaded = extract(Document(path))
        validate_plan(path, style_id, facts, quotes, reloaded)
        manifest.append(record(f"plans/{filename}", "plan", style_id, facts, quotes))

    for filename, style_id, builder, has_contract in CONTRACTS_SPEC:
        doc, quotes = builder()
        facts = empty_contract_facts(has_contract)
        path = CONTRACTS / filename
        doc.save(path)
        if not zipfile.is_zipfile(path):
            raise SystemExit(f"not a docx: {path}")
        with zipfile.ZipFile(path) as zf:
            if "word/document.xml" not in zf.namelist():
                raise SystemExit(f"missing document.xml: {path}")
        reloaded = extract(Document(path))
        validate_contract(path, style_id, facts, quotes, reloaded)
        manifest.append(record(f"contracts/{filename}", "contract", style_id, facts, quotes))

    # Cross-plan coverage.
    plan_items = [item for item in manifest if item["kind"] == "plan"]
    if len(plan_items) != 12 or len(manifest) != 20:
        raise SystemExit("file count")
    for key in ("revenue_monthly_rub", "costs_monthly_rub", "own_funds_rub",
                "months_to_first_cashflow", "use_of_funds", "context_specific"):
        values = [item[key] for item in plan_items]
        if any(v is not None for v in values) and any(v is None for v in values):
            continue
        raise SystemExit(f"fact {key} is not both filled and empty")
    missing_core = 0
    for item in plan_items:
        if item["revenue_monthly_rub"] is None or item["costs_monthly_rub"] is None or item["use_of_funds"] is None:
            missing_core += 1
    if missing_core < 3:
        raise SystemExit("not enough plans missing revenue, costs, or purpose")
    if sum(item["context_specific"] is False for item in plan_items) < 2:
        raise SystemExit("not enough general contexts")
    if sum(item["context_specific"] is True for item in plan_items) < 2:
        raise SystemExit("not enough specific contexts")
    ratios = {
        220000: (400000, 55),
        72000: (80000, 90),
        500000: (400000, 125),
        825000: (1500000, 55),
        1875000: (1500000, 125),
        44000: (80000, 55),
        360000: (400000, 90),
    }
    seen_costs = set()
    for item in plan_items:
        cost = item["costs_monthly_rub"]
        if cost is None:
            continue
        seen_costs.add(cost)
        revenue, num = ratios[cost]
        if item["revenue_monthly_rub"] != revenue or (revenue * num) // 100 != cost:
            raise SystemExit(f"bad cost pair {item['file']}")
    if seen_costs != set(ratios):
        raise SystemExit(f"cost coverage {seen_costs}")

    months = {item["months_to_first_cashflow"] for item in plan_items}
    if not {0, 2, 6, 10}.issubset(months):
        raise SystemExit(f"months coverage {months}")
    purposes = {item["use_of_funds"] for item in plan_items}
    if not {"запуск", "оборудование", "оборотные средства", None}.issubset(purposes):
        raise SystemExit(f"purpose coverage {purposes}")

    out = ROOT / "manifest.json"
    out.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {len(manifest)} files and {out}")
    for item in plan_items:
        print(
            item["style_id"],
            "rev", item["revenue_monthly_rub"],
            "cost", item["costs_monthly_rub"],
            "own", item["own_funds_rub"],
            "m", item["months_to_first_cashflow"],
            "use", item["use_of_funds"],
            "ctx", item["context_specific"],
        )


if __name__ == "__main__":
    main()
